"""Xuất heatmap: bảng text, CSV, Excel."""
from __future__ import annotations

import csv
from pathlib import Path

from .analyzer import DayResult
from .config import Config


def _columns(days: list[DayResult], cfg: Config) -> list[str]:
    if cfg.columns:
        return cfg.columns
    keys = {k for d in days for k in d.first}
    return sorted(keys)


def _best(day: DayResult, cfg: Config) -> str | None:
    pool = cfg.highlight or list(day.first)
    cand = [(day.first[k].delay_ms, k) for k in pool
            if k in day.first and day.first[k].delay_ms is not None]
    return min(cand)[1] if cand else None


def _fmt(day: DayResult, key: str) -> str:
    a = day.first.get(key)
    return "-" if a is None or a.delay_ms is None else f"{a.delay_ms:.2f}"


def render_table(days: list[DayResult], cfg: Config) -> str:
    cols = _columns(days, cfg)
    rows = [["Date", *cols]]
    for d in days:
        best = _best(d, cfg)
        rows.append([d.date] + [_fmt(d, k) + ("*" if k == best else "") for k in cols])
    widths = [max(len(r[i]) for r in rows) for i in range(len(rows[0]))]
    line = lambda r: "  ".join(c.rjust(w) if i else c.ljust(w) for i, (c, w) in enumerate(zip(r, widths)))
    out = [line(rows[0]), "  ".join("-" * w for w in widths)]
    out += [line(r) for r in rows[1:]]
    out.append("\n(*) = message tới sớm nhất trong nhóm highlight; đơn vị ms")
    return "\n".join(out)


def render_race(days: list[DayResult]) -> str:
    out = []
    for d in days:
        seq = " -> ".join(f"{k} (+{gap:.2f}ms)" for k, gap in d.race())
        out.append(f"{d.date}: {seq or '(không có message)'}")
    return "\n".join(out)


def write_csv(days: list[DayResult], cfg: Config, path: str | Path) -> None:
    cols = _columns(days, cfg)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["Date", *cols, "Noted"])
        for d in days:
            w.writerow([d.date, *[_fmt(d, k).replace("-", "") for k in cols], cfg.notes.get(d.date, "")])


def write_xlsx(days: list[DayResult], cfg: Config, path: str | Path) -> None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
    except ImportError as e:  # pragma: no cover
        raise SystemExit("Cần cài openpyxl: pip install openpyxl") from e
    cols = _columns(days, cfg)
    wb = Workbook()
    ws = wb.active
    ws.title = "Latency Heatmap"
    ws.append(["Date", *[cfg.labels.get(k, f"First {k}") for k in cols], "Noted"])
    for c in ws[1]:
        c.font = Font(bold=True)
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    green, red = PatternFill("solid", fgColor="E3F0D8"), Font(bold=True, color="C00000")
    for r, d in enumerate(days, start=2):
        best = _best(d, cfg)
        ws.cell(r, 1, d.date)
        for i, k in enumerate(cols, start=2):
            a = d.first.get(k)
            cell = ws.cell(r, i, round(a.delay_ms, 2) if a and a.delay_ms is not None else None)
            cell.number_format = "0.00"
            if k == best:
                cell.fill, cell.font = green, red
        ws.cell(r, len(cols) + 2, cfg.notes.get(d.date, ""))
    ws.column_dimensions["A"].width = 12
    wb.save(path)


def table_data(days: list[DayResult], cfg: Config) -> dict:
    """Dữ liệu heatmap dạng JSON cho web UI."""
    cols = _columns(days, cfg)
    rows = []
    for d in days:
        rows.append({
            "date": d.date,
            "cells": {k: (round(d.first[k].delay_ms, 2) if k in d.first and d.first[k].delay_ms is not None else None)
                      for k in cols},
            "best": _best(d, cfg),
            "note": cfg.notes.get(d.date, ""),
            "packets": d.packets,
        })
    return {"columns": cols, "labels": {k: cfg.labels.get(k, f"First {k}") for k in cols}, "rows": rows}
